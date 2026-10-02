"""Classeur de restitution des résultats : S.25.01, modules, MCR, best estimates."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from scr_data.export import POLICE, TOTAL, _entete, _style, _titre


def _ecrire_table(ws, df: pd.DataFrame, entetes: list[str], formats: dict[int, str],
                  lignes_totales: tuple[str, ...] = ()) -> None:
    _entete(ws, 3, entetes)
    for i, r in enumerate(df.itertuples(index=False), start=4):
        for j, v in enumerate(r, start=1):
            cellule = ws.cell(row=i, column=j, value=v)
            if j in formats:
                cellule.number_format = formats[j]
        libelle = str(r[0])
        if any(mot in libelle for mot in lignes_totales):
            for j in range(1, len(entetes) + 1):
                ws.cell(row=i, column=j).fill = TOTAL
                ws.cell(row=i, column=j).font = Font(name=POLICE, bold=True)


def ecrire_rapport(chemin: Path, hyp: dict[str, Any], resultats: dict[str, Any]) -> None:
    wb = Workbook()
    nom, date = hyp["general"]["nom_compagnie"], hyp["general"]["date_evaluation"]
    euro = "#,##0.0"

    ws = wb.active
    ws.title = "S.25.01"
    _titre(ws, f"Capital de solvabilité requis — {nom}",
           f"Formule standard · date d'évaluation {date} · montants en M€")
    _ecrire_table(ws, resultats["s25"], ["Code", "Poste", "Montant (M€)"], {3: euro},
                  ("BSCR", "Capital de solvabilité requis (SCR)"))
    ligne = 4 + len(resultats["s25"]) + 1
    bof = resultats["moteur"].bof_base
    for libelle, valeur, fmt in (("Fonds propres de base", bof, euro),
                                 ("Ratio de couverture du SCR", bof / resultats["scr"], "0%"),
                                 ("MCR", resultats["mcr_valeur"], euro),
                                 ("Ratio de couverture du MCR", bof / resultats["mcr_valeur"], "0%")):
        ws.cell(row=ligne, column=2, value=libelle).font = Font(name=POLICE, bold=True)
        c = ws.cell(row=ligne, column=3, value=valeur)
        c.number_format = fmt
        c.font = Font(name=POLICE, bold=True)
        ligne += 1
    _style(ws, {"A": 10, "B": 56, "C": 16})

    for onglet, cle, entetes in (("Marche", "marche", ["Sous-module", "SCR (M€)", "Détail"]),
                                 ("Vie", "vie", ["Sous-module", "SCR (M€)", "Détail"]),
                                 ("Sante", "sante", ["Sous-module", "SCR (M€)"]),
                                 ("NonVie", "non_vie", ["Sous-module", "SCR (M€)", "Détail"]),
                                 ("Contrepartie", "contrepartie", ["Composante", "SCR (M€)", "Détail"])):
        ws = wb.create_sheet(onglet)
        _titre(ws, f"Détail — {onglet}", "Chaque sous-module est évalué en reprojetant les passifs")
        _ecrire_table(ws, resultats[cle], entetes, {2: euro}, ("agrégé",))
        _style(ws, {"A": 34, "B": 14, "C": 52})

    ws = wb.create_sheet("BSCR")
    _titre(ws, "Agrégation des modules", "Matrice de corrélation de l'annexe IV")
    _ecrire_table(ws, resultats["bscr"], ["Module", "SCR (M€)", "Part du BSCR"], {2: euro, 3: "0.0%"},
                  ("BSCR",))
    _style(ws, {"A": 22, "B": 14, "C": 14})

    ws = wb.create_sheet("MCR")
    _titre(ws, "Minimum de capital requis", "Calcul linéaire, corridor 25 % – 45 % du SCR")
    _ecrire_table(ws, resultats["mcr"], ["Poste", "Montant (M€)"], {2: euro}, ("MCR retenu",))
    _style(ws, {"A": 38, "B": 16})

    ws = wb.create_sheet("BestEstimate")
    _titre(ws, "Best estimate par segment", "Projection déterministe, sans TVOG (V1)")
    _ecrire_table(ws, resultats["be"], ["Segment", "BE (M€)"], {2: euro}, ("total",))
    _style(ws, {"A": 30, "B": 16})

    wb.save(chemin)
