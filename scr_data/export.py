"""Export des données : un CSV par table + un classeur Excel de synthèse."""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

POLICE = "Arial"
ENTETE = PatternFill("solid", start_color="1F3864")
TOTAL = PatternFill("solid", start_color="D9E1F2")
FIN = Side(style="thin", color="BFBFBF")


def exporter_csv(tables: dict[str, pd.DataFrame], dossier: Path) -> None:
    dossier.mkdir(parents=True, exist_ok=True)
    for nom, df in tables.items():
        df.to_csv(dossier / f"{nom}.csv", index=False, sep=";", decimal=",", encoding="utf-8-sig")


def _entete(ws, ligne: int, valeurs: list[str]) -> None:
    for j, v in enumerate(valeurs, start=1):
        c = ws.cell(row=ligne, column=j, value=v)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = ENTETE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style(ws, largeurs: dict[str, int]) -> None:
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != POLICE:
                c.font = Font(name=POLICE, bold=c.font.bold, color=c.font.color, italic=c.font.italic)
    for col, l in largeurs.items():
        ws.column_dimensions[col].width = l
    ws.freeze_panes = "A4"


def _titre(ws, titre: str, sous_titre: str) -> None:
    ws["A1"] = titre
    ws["A1"].font = Font(name=POLICE, bold=True, size=14, color="1F3864")
    ws["A2"] = sous_titre
    ws["A2"].font = Font(name=POLICE, italic=True, size=9, color="595959")


def exporter_excel(chemin: Path, hyp: dict[str, Any], bilan: pd.DataFrame, controles: pd.DataFrame,
                   inventaire: pd.DataFrame, courbe: pd.DataFrame, info_courbe: dict[str, Any]) -> None:
    wb = Workbook()
    nom = hyp["general"]["nom_compagnie"]
    date = hyp["general"]["date_evaluation"]

    # --- Bilan -------------------------------------------------------------
    ws = wb.active
    ws.title = "Bilan"
    _titre(ws, f"Bilan prudentiel indicatif — {nom}", f"Date d'évaluation {date} · montants en M€ · "
           "les BE marqués « cible à recalculer » seront remplacés par le modèle de projection")
    _entete(ws, 3, ["Côté", "Code S.02.01", "Poste", "Montant (M€)", "Statut"])
    ligne = 4
    lignes_total = {}
    for cote in ("Actif", "Passif"):
        debut = ligne
        for _, r in bilan[bilan["cote"] == cote].iterrows():
            valeurs = [r["cote"], r["code"], r["poste"], round(float(r["montant"]), 3), r["statut"]]
            for j, v in enumerate(valeurs, start=1):
                ws.cell(row=ligne, column=j, value=v)
            ws.cell(row=ligne, column=4).number_format = "#,##0.0"
            ligne += 1
        ws.cell(row=ligne, column=3, value=f"Total {cote.lower()}")
        ws.cell(row=ligne, column=4, value=f"=SUM(D{debut}:D{ligne - 1})")
        for j in range(1, 6):
            ws.cell(row=ligne, column=j).fill = TOTAL
            ws.cell(row=ligne, column=j).font = Font(name=POLICE, bold=True)
        ws.cell(row=ligne, column=4).number_format = "#,##0.0"
        lignes_total[cote] = ligne
        ligne += 2
    ws.cell(row=ligne, column=3, value="Fonds propres (actif − passif)").font = Font(name=POLICE, bold=True)
    c = ws.cell(row=ligne, column=4, value=f"=D{lignes_total['Actif']}-D{lignes_total['Passif']}")
    c.number_format = "#,##0.0"
    c.font = Font(name=POLICE, bold=True)
    ws.cell(row=ligne + 1, column=3, value="Fonds propres / total actif")
    c = ws.cell(row=ligne + 1, column=4, value=f"=IFERROR(D{ligne}/D{lignes_total['Actif']},0)")
    c.number_format = "0.0%"
    _style(ws, {"A": 9, "B": 13, "C": 64, "D": 15, "E": 22})

    # --- Allocation --------------------------------------------------------
    ws = wb.create_sheet("Allocation")
    _titre(ws, "Allocation des placements hors UC (après transparence)",
           "Valeurs de marché issues de l'inventaire ligne par ligne")
    _entete(ws, 3, ["Classe", "Sous-classe", "Valeur de marché (M€)", "Poids", "Nb lignes"])
    euro = inventaire[inventaire["portefeuille"] == "euro"]
    agg = (euro.groupby(["classe", "sous_classe"])
           .agg(vm=("valeur_marche", "sum"), n=("id_ligne", "count")).reset_index())
    debut = 4
    fin = debut + len(agg) - 1
    for i, r in enumerate(agg.itertuples(), start=debut):
        ws.append([r.classe, r.sous_classe, round(r.vm, 3), f"=IFERROR(C{i}/C${fin + 1},0)", r.n])
        ws.cell(row=i, column=3).number_format = "#,##0.0"
        ws.cell(row=i, column=4).number_format = "0.0%"
    ws.append(["Total", "", f"=SUM(C{debut}:C{fin})", f"=SUM(D{debut}:D{fin})", f"=SUM(E{debut}:E{fin})"])
    for j in range(1, 6):
        ws.cell(row=fin + 1, column=j).fill = TOTAL
        ws.cell(row=fin + 1, column=j).font = Font(name=POLICE, bold=True)
    ws.cell(row=fin + 1, column=3).number_format = "#,##0.0"
    ws.cell(row=fin + 1, column=4).number_format = "0.0%"
    _style(ws, {"A": 22, "B": 28, "C": 22, "D": 10, "E": 11})

    # --- Contrôles ---------------------------------------------------------
    ws = wb.create_sheet("Controles")
    _titre(ws, "Contrôles de calibrage", "OK : écart dans la tolérance · ALERTE : à investiguer · INFO : indicateur")
    _entete(ws, 3, ["Contrôle", "Valeur", "Cible", "Écart relatif", "Statut"])
    for i, r in enumerate(controles.itertuples(), start=4):
        ws.append([r.controle, float(r.valeur), None if pd.isna(r.cible) else float(r.cible),
                   None if pd.isna(r.ecart_relatif) else float(r.ecart_relatif), r.statut])
        ws.cell(row=i, column=2).number_format = "#,##0.000"
        ws.cell(row=i, column=3).number_format = "#,##0.000"
        ws.cell(row=i, column=4).number_format = "0.00%"
        couleur = {"OK": "548235", "ALERTE": "C00000", "INFO": "2F5597"}[r.statut]
        ws.cell(row=i, column=5).font = Font(name=POLICE, bold=True, color=couleur)
    _style(ws, {"A": 70, "B": 14, "C": 14, "D": 14, "E": 10})

    # --- Courbe ------------------------------------------------------------
    ws = wb.create_sheet("Courbe")
    alpha = info_courbe["alpha"]
    _titre(ws, "Courbe des taux sans risque", f"Source : {info_courbe['source']}"
           + (f" · alpha = {alpha:.4f}" if alpha else "") + f" · UFR = {hyp['courbe']['ufr']:.2%}")
    _entete(ws, 3, ["Maturité", "Taux spot", "Facteur d'actualisation", "Forward 1 an"])
    for i, r in enumerate(courbe.itertuples(), start=4):
        ws.append([int(r.maturite), float(r.taux_spot), float(r.facteur_actualisation), float(r.forward_1an)])
        ws.cell(row=i, column=2).number_format = "0.000%"
        ws.cell(row=i, column=3).number_format = "0.000000"
        ws.cell(row=i, column=4).number_format = "0.000%"
    _style(ws, {"A": 11, "B": 13, "C": 22, "D": 14})

    for ws in wb.worksheets:
        for col in range(1, ws.max_column + 1):
            ws.cell(row=3, column=col).border = Border(bottom=FIN)
    wb.save(chemin)
